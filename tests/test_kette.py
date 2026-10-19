# -*- coding: utf-8 -*-
"""End-to-End-Test: python tests/test_kette.py  (benötigt GDAL, numpy, openpyxl, scipy)"""
import os
import sys
import tempfile

import numpy as np
from osgeo import gdal, ogr

HIER = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HIER))
sys.path.insert(0, HIER)

import dgm1945_kern as k  # noqa: E402
import testdaten_erzeugen as td  # noqa: E402


def lies(p, quelle=None):
    a = gdal.Open(p).ReadAsArray().astype(float)
    ok = a != k.NODATA
    if quelle:
        ok &= gdal.Open(quelle).ReadAsArray() == 1
    return a[ok]


def lauf(ordner, **kw):
    e = k.einstellungen_aus_ini(os.path.join(ordner, "test.ini"))
    for n, v in kw.items():
        setattr(e, n, v)
    return k.ausfuehren(e, log=lambda m: None)


def main():
    with tempfile.TemporaryDirectory() as d:
        td.main(d)

        # 1) Standard: Excel + Lot + DGM1-Kacheln
        r = lauf(d)
        diff = lies(r["differenz"], r["quelle"])
        assert abs(np.median(diff) - 1.0) < 0.01 and diff.min() > 0.8, diff
        assert np.allclose(lies(r["gok45"]) - 8, lies(r["eindringtiefe"]), atol=1e-4)
        assert abs(np.median(lies(r["sondiertiefe"], r["quelle"])) - 9.0) < 0.01
        aussen = lies(r["differenz"])
        assert np.sum(np.abs(aussen) < 1e-6) > 0, "DGM1-Auffüllung fehlt"
        print("Standard ok (inkl. DGM1-Auffüllung)")

        # 2) CSV-Eingabe mit Komma-Dezimal
        import openpyxl
        ws = openpyxl.load_workbook(os.path.join(d, "hoehen_1945.xlsx")).active
        with open(os.path.join(d, "h.csv"), "w", encoding="utf-8") as f:
            for row in list(ws.iter_rows(values_only=True))[2:]:
                f.write(";".join("" if v is None else str(v).replace(".", ",") if isinstance(v, float)
                                 else str(v) for v in row) + "\n")
        r = lauf(d, excel=os.path.join(d, "h.csv"), ausgabe_ordner=os.path.join(d, "e_csv"))
        assert abs(np.median(lies(r["differenz"], r["quelle"])) - 1.0) < 0.01
        print("CSV ok")

        # 3) km-Kalibrierpunkte statt km_start
        ds = ogr.GetDriverByName("GPKG").CreateDataSource(os.path.join(d, "km.gpkg"))
        lds = ogr.Open(os.path.join(d, "linien.shp"))
        lyr = ds.CreateLayer("km", lds.GetLayer().GetSpatialRef(), ogr.wkbPoint)
        lyr.CreateField(ogr.FieldDefn("km", ogr.OFTReal))
        ach = k.Polylinie(td.achse_xy())
        for m in (0, 1000, 2000):
            f = ogr.Feature(lyr.GetLayerDefn())
            f.SetField("km", 10 + m / 1000)
            g = ogr.Geometry(ogr.wkbPoint)
            g.AddPoint_2D(*map(float, ach.punkt(m)))
            f.SetGeometry(g)
            lyr.CreateFeature(f)
        ds = None
        r = lauf(d, km_start=None, km_punkte=os.path.join(d, "km.gpkg"),
                 ausgabe_ordner=os.path.join(d, "e_km"))
        assert abs(np.median(lies(r["differenz"], r["quelle"])) - 1.0) < 0.01
        print("km-Punkte ok")

        # 4) Excel ohne openpyxl (GDAL-Fallback), Kopfzeile in Zeile 1 nötig
        wb = openpyxl.Workbook()
        for row in list(ws.iter_rows(values_only=True))[2:]:
            wb.active.append(row)
        wb.save(os.path.join(d, "h1.xlsx"))
        sys.modules["openpyxl"] = None
        r = lauf(d, excel=os.path.join(d, "h1.xlsx"), ausgabe_ordner=os.path.join(d, "e_gdal"))
        del sys.modules["openpyxl"]
        assert abs(np.median(lies(r["differenz"], r["quelle"])) - 1.0) < 0.01
        print("Excel via GDAL ok")

        # 5) Methode "entlang" läuft durch
        r = lauf(d, methode="entlang", ausgabe_ordner=os.path.join(d, "e_ent"))
        assert lies(r["gok45"]).size > 0
        print("Methode entlang ok")

        # 6) Station + Seite + Abstand (Tabellenaufbau KILOMETRIE;SEITE;ABSTAND;HOEHE),
        #    nur Achse mit DBF-Feld START_M, km fallend, doppelte Planjahre
        drv = ogr.GetDriverByName("ESRI Shapefile")
        ads = drv.CreateDataSource(os.path.join(d, "achse.shp"))
        alyr = ads.CreateLayer("achse", lds.GetLayer().GetSpatialRef(), ogr.wkbLineString)
        alyr.CreateField(ogr.FieldDefn("START_M", ogr.OFTReal))
        f = ogr.Feature(alyr.GetLayerDefn())
        f.SetField("START_M", 23670.0)
        g = ogr.Geometry(ogr.wkbLineString)
        for x, y in td.achse_xy():
            g.AddPoint_2D(float(x), float(y))
        f.SetGeometry(g)
        alyr.CreateFeature(f)
        ads = None
        # km fallend in Digitalisierrichtung: Station s (Achslänge) -> 23670 - s
        # Bezugsrichtung = steigende km = gegen Digitalisierrichtung -> "links" = rechts der Digitalisierung
        with open(os.path.join(d, "abstand.csv"), "w", encoding="cp1252") as fh:
            fh.write("KILOMETRIE;SEITE;ABSTAND;HOEHE;JAHR;BEMERKUNG\n")
            for s_ in range(0, 2001, 100):
                stn = 23670 - s_
                for typ, ab in (("Sohle", 0), ("Sohle", 20), ("Berme", 35), ("Deich", 50)):
                    for seite in (("links", "rechts") if ab else ("links",)):
                        dig_seite = "rechts" if seite == "links" else "links"
                        hh = td.soll_hoehe(typ, s_, dig_seite)
                        fh.write(f"{stn};{seite};{ab};{str(round(hh, 3)).replace('.', ',')};1949;Höhe ok\n")
                        # späterer Plan am selben Ort mit falscher Höhe -> muss verworfen werden
                        fh.write(f"{stn};{seite};{ab};{str(round(hh + 3, 3)).replace('.', ',')};1960;\n")
        r = lauf(d, excel=os.path.join(d, "abstand.csv"), linien=os.path.join(d, "achse.shp"),
                 km_start=None, km_steigend=None, ausgabe_ordner=os.path.join(d, "e_abst"))
        diff = lies(r["differenz"], r["quelle"])
        assert abs(np.median(diff) - 1.0) < 0.02 and np.all(np.abs(np.percentile(diff, [1, 99]) - 1) < 0.05), np.percentile(diff, [1, 50, 99])
        print("Station/Seite/Abstand + DBF-Start + Jahr-Auswahl ok")
    print("ALLE TESTS OK")


if __name__ == "__main__":
    main()
